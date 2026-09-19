"""Medical report analysis: extract lab values from an uploaded document
(PDF, scanned PDF, image, DOCX, or text), compare against reference ranges,
and produce a plain-language report.

Extraction is defensive and multi-format:
  - PDF: embedded text via pdfplumber; scanned PDFs fall back to OCR.
  - Images: OCR with preprocessing (grayscale, upscale, autocontrast) and
    multiple page-segmentation modes; the best result is kept.
  - DOCX: paragraphs + tables via python-docx.
  - TXT / CSV / plain text: decoded directly.

Parsing: per-marker alias + numeric regex, with reference-range detection,
H/L flag detection, and guards against confusing related markers.
No LLM required.
"""
from __future__ import annotations

import io
import re
from typing import Any

from app.services.lab_reference import MARKERS

try:
    import pdfplumber
except Exception:  # pragma: no cover
    pdfplumber = None
try:
    import pymupdf  # PyMuPDF
except Exception:  # pragma: no cover
    pymupdf = None
try:
    import pytesseract
    from PIL import Image, ImageOps, ImageFilter
except Exception:  # pragma: no cover
    pytesseract = None
    Image = ImageOps = ImageFilter = None
try:
    import docx  # python-docx
except Exception:  # pragma: no cover
    docx = None


class ReportError(Exception):
    pass


# --------------------------------------------------------------------------- #
# Extraction
# --------------------------------------------------------------------------- #
_OCR_CONFIGS = ["--oem 3 --psm 6", "--oem 3 --psm 4", "--oem 3 --psm 3", "--oem 3 --psm 11"]


def _prep_image(img: "Image.Image") -> "Image.Image":
    """Grayscale, upscale small images, and boost contrast for better OCR."""
    try:
        img = img.convert("L")
        w, h = img.size
        longest = max(w, h)
        if longest < 1800:
            scale = 1800.0 / longest
            img = img.resize((int(w * scale), int(h * scale)))
        img = ImageOps.autocontrast(img)
        img = img.filter(ImageFilter.SHARPEN)
    except Exception:
        pass
    return img


def _ocr_image(img: "Image.Image") -> str:
    if pytesseract is None:
        return ""
    prepped = _prep_image(img)
    best = ""
    for cfg in _OCR_CONFIGS:
        try:
            txt = pytesseract.image_to_string(prepped, config=cfg)
        except Exception:
            txt = ""
        if len(txt.strip()) > len(best.strip()):
            best = txt
    return best


def _ocr_image_bytes(data: bytes) -> str:
    if Image is None:
        return ""
    try:
        return _ocr_image(Image.open(io.BytesIO(data)))
    except Exception:
        return ""


def _ocr_pdf(data: bytes) -> str:
    if pymupdf is None or pytesseract is None:
        return ""
    out = []
    try:
        doc = pymupdf.open(stream=data, filetype="pdf")
        for page in doc:
            pix = page.get_pixmap(dpi=300)
            out.append(_ocr_image(Image.open(io.BytesIO(pix.tobytes("png")))))
    except Exception:
        return ""
    return "\n".join(out)


def _read_docx(data: bytes) -> str:
    if docx is None:
        return ""
    try:
        d = docx.Document(io.BytesIO(data))
        parts = [p.text for p in d.paragraphs]
        for table in d.tables:
            for row in table.rows:
                parts.append(" ".join(c.text for c in row.cells))
        return "\n".join(parts)
    except Exception:
        return ""


def extract_text(data: bytes, filename: str, content_type: str) -> tuple[str, str]:
    """Return (text, method). Raises ReportError with a friendly message."""
    name = (filename or "").lower()
    ctype = (content_type or "").lower()

    def ext(*e):
        return any(name.endswith(x) for x in e)

    # PDF
    if ext(".pdf") or "pdf" in ctype:
        text = ""
        if pdfplumber is not None:
            try:
                with pdfplumber.open(io.BytesIO(data)) as pdf:
                    text = "\n".join((p.extract_text() or "") for p in pdf.pages)
            except Exception:
                text = ""
        if len(text.strip()) >= 40:
            return text, "pdf-text"
        ocr = _ocr_pdf(data)
        if ocr.strip():
            return ocr, "pdf-ocr"
        raise ReportError("Could not read text from this PDF. If it is a scan, try a clearer copy.")

    # DOCX
    if ext(".docx") or "officedocument.wordprocessing" in ctype:
        text = _read_docx(data)
        if text.strip():
            return text, "docx"
        raise ReportError("Could not read this Word document. Try exporting it as a PDF.")

    # Plain text / CSV
    if ext(".txt", ".csv", ".tsv") or ctype.startswith("text/"):
        for enc in ("utf-8", "latin-1"):
            try:
                return data.decode(enc), "text"
            except Exception:
                continue
        raise ReportError("Could not decode this text file.")

    # Images
    if ext(".png", ".jpg", ".jpeg", ".webp", ".bmp", ".tif", ".tiff", ".gif", ".heic") or ctype.startswith("image/"):
        ocr = _ocr_image_bytes(data)
        if ocr.strip():
            return ocr, "image-ocr"
        raise ReportError("Could not read text from this image. Try a clearer, well-lit, higher-resolution photo.")

    # Last resort: try text, then treat as image
    for enc in ("utf-8",):
        try:
            t = data.decode(enc)
            if t.strip():
                return t, "text"
        except Exception:
            pass
    ocr = _ocr_image_bytes(data)
    if ocr.strip():
        return ocr, "image-ocr"
    raise ReportError("Unsupported or unreadable file. Upload a PDF, image, Word document, or text file.")


# --------------------------------------------------------------------------- #
# Parsing
# --------------------------------------------------------------------------- #
_NUM = r"([0-9]{1,3}(?:,[0-9]{3})*(?:\.[0-9]+)?|[0-9]+(?:\.[0-9]+)?)"
_RANGE_RE = re.compile(rf"{_NUM}\s*(?:[-–—]|to)\s*{_NUM}")
_UPPER_RE = re.compile(rf"[<≤]\s*{_NUM}")   # "< 200"
_LOWER_RE = re.compile(rf"[>≥]\s*{_NUM}")   # "> 40"
_FLAG_RE = re.compile(r"\b(H|HIGH|HI|L|LOW|LO)\b")


def _clean_num(s: str) -> float:
    return float(s.replace(",", ""))


def _find_marker(text_lc: str, marker: dict) -> dict | None:
    for alias in sorted(marker["aliases"], key=len, reverse=True):
        pattern = re.compile(re.escape(alias) + r"\b", re.IGNORECASE)
        for m in pattern.finditer(text_lc):
            start, end = m.start(), m.end()
            prefix = text_lc[max(0, start - 14):start]
            if any(x in prefix for x in marker.get("exclude_prefix", [])):
                continue
            window = text_lc[end:end + 70]
            num_match = re.search(_NUM, window)
            if not num_match:
                continue
            value = _clean_num(num_match.group(1))
            if value <= 0 or value > 100000:
                continue

            after = window[num_match.end():num_match.end() + 45]
            ref_low, ref_high, ref_source = marker["low"], marker["high"], "standard"
            rng = _RANGE_RE.search(after)
            if rng:
                lo, hi = _clean_num(rng.group(1)), _clean_num(rng.group(2))
                if hi > lo and hi < 100000:
                    ref_low, ref_high, ref_source = lo, hi, "report"
            else:
                up = _UPPER_RE.search(after)
                lw = _LOWER_RE.search(after)
                if up:
                    ref_high, ref_source = _clean_num(up.group(1)), "report"
                elif lw:
                    ref_low, ref_source = _clean_num(lw.group(1)), "report"

            # explicit H/L flag near the value (only used when no range detected)
            flag = None
            if ref_source == "standard":
                fm = _FLAG_RE.search(after.upper())
                if fm:
                    flag = "High" if fm.group(1).startswith(("H", "HI")) else "Low"

            return {"value": value, "ref_low": ref_low, "ref_high": ref_high,
                    "ref_source": ref_source, "flag": flag}
    return None


def _status(value: float, low: float, high: float, flag: str | None) -> str:
    if value < low:
        return "Low"
    if value > high:
        return "High"
    return flag or "Normal"


def analyze(data: bytes, filename: str, content_type: str) -> dict[str, Any]:
    text, method = extract_text(data, filename, content_type)
    return analyze_text(text, method)


def analyze_text(text: str, method: str = "text") -> dict[str, Any]:
    text_lc = re.sub(r"[ \t]+", " ", text.lower())

    results: list[dict] = []
    for marker in MARKERS:
        found = _find_marker(text_lc, marker)
        if not found:
            continue
        status = _status(found["value"], found["ref_low"], found["ref_high"], found.get("flag"))
        note = marker["high_note"] if status == "High" else marker["low_note"] if status == "Low" else ""
        results.append({
            "key": marker["key"], "name": marker["name"], "category": marker["category"],
            "value": found["value"], "unit": marker["unit"],
            "ref_low": found["ref_low"], "ref_high": found["ref_high"],
            "ref_source": found["ref_source"], "status": status,
            "about": marker["about"], "note": note,
            "advice": marker["advice"] if status != "Normal" else "",
            "related": marker.get("related", ""),
        })

    abnormal = [r for r in results if r["status"] != "Normal"]
    categories: dict[str, list[dict]] = {}
    for r in results:
        categories.setdefault(r["category"], []).append(r)
    related = sorted({r["related"] for r in abnormal if r["related"]})

    if not results:
        summary = ("We could not confidently detect standard lab values in this document. It may use an "
                   "unusual layout or be a document type this tool does not yet parse. You can still try a "
                   "clearer PDF or a sharper photo.")
    elif not abnormal:
        summary = (f"We detected {len(results)} lab value(s), and all of them fall within their normal "
                   "reference ranges. Keep up your current healthy habits.")
    else:
        names = ", ".join(r["name"] for r in abnormal[:5])
        summary = (f"We detected {len(results)} lab value(s); {len(abnormal)} fall outside the normal "
                   f"reference range: {names}{'…' if len(abnormal) > 5 else ''}. See the flagged items "
                   "below for plain-language explanations and next steps.")

    return {
        "method": method,
        "summary": summary,
        "counts": {"total": len(results), "abnormal": len(abnormal), "normal": len(results) - len(abnormal)},
        "flagged": abnormal,
        "results": results,
        "categories": [{"name": k, "markers": v} for k, v in categories.items()],
        "related_assessments": related,
        "disclaimer": (
            "This analysis is generated automatically from the values detected in your document and compared "
            "to general adult reference ranges, which vary by lab, age, and sex. It is for educational purposes "
            "only, is not a diagnosis, and may miss or misread values. Always review your report with a "
            "qualified healthcare professional."
        ),
    }
